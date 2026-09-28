from __future__ import annotations

import json
import re
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
import polars as pl

from opfl.constants import TEAM_ABBREV_NORMALIZE
from opfl.data_fetcher import normalize_name
from opfl.excel_parser import parse_player_name
from opfl.scorer import OPFLScorer


SOURCE = Path("/Users/griffin/Downloads/OPFL Scoring 2026.xlsx")
OUTPUT = Path("outputs/opfl_week2_2026/OPFL Scoring 2026 - Week 2 scored.xlsx")
SEASON = 2026
WEEK = 2

POSITION_ROWS = {
    1: {
        "QB": range(2, 5),
        "RB": range(6, 10),
        "WR": range(11, 16),
        "TE": range(17, 20),
        "K": range(21, 23),
        "DF": range(24, 26),
        "HC": range(27, 29),
    },
    39: {
        "QB": range(40, 43),
        "RB": range(44, 48),
        "WR": range(49, 54),
        "TE": range(55, 58),
        "K": range(59, 61),
        "DF": range(62, 64),
        "HC": range(65, 67),
    },
}
POSITION_ORDER = ["QB", "RB", "WR", "TE", "K", "DF", "HC"]
EXPECTED_STARTERS = {"QB": 1, "RB": 2, "WR": 2, "TE": 1, "K": 1, "DF": 1, "HC": 1}
SCORING_OFFSETS = [0, 1, 2, 3, 4, 5, 7, 9, 11]
PLAYER_COLUMNS = [4, 7, 10, 13, 16, 19]


@dataclass
class Starter:
    position: str
    roster_row: int
    roster_value: str
    name: str
    nfl_team: str
    score: float | None = None
    raw_inputs: list[float | int | None] = field(default_factory=list)
    found: bool = False
    matched_name: str = ""
    pending: bool = False


@dataclass
class FantasyTeam:
    index: int
    roster_header: str
    display_name: str
    player_col: int
    header_row: int
    starters: list[Starter]

    @property
    def total(self) -> float:
        return float(sum(starter.score or 0 for starter in self.starters))

    @property
    def pending_count(self) -> int:
        return sum(starter.pending for starter in self.starters)


def normalized_team(team: str) -> str:
    return TEAM_ABBREV_NORMALIZE.get(team.upper(), team.upper())


def clean_header(value: Any) -> str:
    return re.sub(r"\s*\(\d+\)\s*$", "", str(value or "")).strip()


def excel_number(value: float | int | None) -> float | int | None:
    if value is None:
        return None
    numeric = float(value)
    return int(numeric) if numeric.is_integer() else numeric


def copy_cell(source, target, value: Any) -> None:
    target.value = value
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)
    target.comment = copy(source.comment)
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)


def read_teams(workbook) -> list[FantasyTeam]:
    roster = workbook["Rosters"]
    scoring = workbook["Scoring"]
    teams: list[FantasyTeam] = []
    team_index = 0

    for header_row in (1, 39):
        for player_col in PLAYER_COLUMNS:
            starters: list[Starter] = []
            by_position: dict[str, int] = {}
            for position in POSITION_ORDER:
                for row in POSITION_ROWS[header_row][position]:
                    if roster.cell(row, player_col - 1).value != "*":
                        continue
                    value = str(roster.cell(row, player_col).value or "").strip()
                    name, nfl_team = parse_player_name(value)
                    starters.append(
                        Starter(
                            position=position,
                            roster_row=row,
                            roster_value=value,
                            name=name,
                            nfl_team=nfl_team,
                        )
                    )
                    by_position[position] = by_position.get(position, 0) + 1

            if by_position != EXPECTED_STARTERS:
                raise ValueError(
                    f"Invalid starters for {roster.cell(header_row, player_col).value}: {by_position}"
                )

            block_start = 2 + 12 * team_index
            teams.append(
                FantasyTeam(
                    index=team_index,
                    roster_header=str(roster.cell(header_row, player_col).value or ""),
                    display_name=str(scoring.cell(block_start, 1).value or "").strip(),
                    player_col=player_col,
                    header_row=header_row,
                    starters=starters,
                )
            )
            team_index += 1

    if len(teams) != 12:
        raise ValueError(f"Expected 12 teams, found {len(teams)}")
    return teams


def game_complete(data, team: str) -> bool:
    team_code = normalized_team(team)
    schedule = data.schedules
    home = schedule.filter(pl.col("home_team") == team_code)
    if home.height:
        return home.row(0, named=True).get("home_score") is not None
    away = schedule.filter(pl.col("away_team") == team_code)
    if away.height:
        return away.row(0, named=True).get("away_score") is not None
    raise ValueError(f"No Week {WEEK} game found for {team_code}")


def player_stats(data, starter: Starter) -> dict[str, Any] | None:
    stats = data.find_player(starter.name, starter.nfl_team, starter.position)
    if not stats:
        return None
    matched = str(stats.get("player_display_name") or "")
    query_norm = normalize_name(starter.name)
    matched_norm = normalize_name(matched)
    if query_norm == matched_norm:
        return stats
    query_last = query_norm.split()[-1] if query_norm else ""
    matched_last = matched_norm.split()[-1] if matched_norm else ""
    if query_last and query_last == matched_last and (
        query_norm in matched_norm or matched_norm in query_norm
    ):
        return stats
    if query_last == matched_last and starter.position in {"QB", "RB", "WR", "TE", "K"}:
        return stats
    raise ValueError(f"Unsafe player match: {starter.name} -> {matched}")


def offensive_raw(data, starter: Starter, stats: dict[str, Any] | None) -> list[int]:
    if stats is None:
        return [0] * 10
    player_id = stats.get("player_id")
    turnover_tds = data.get_turnovers_returned_for_td(player_id) if player_id else {}
    touchdowns = sum(
        int(stats.get(field, 0) or 0)
        for field in ("passing_tds", "rushing_tds", "receiving_tds", "fumble_recovery_tds")
    )
    two_points = sum(
        int(stats.get(field, 0) or 0)
        for field in (
            "passing_2pt_conversions",
            "rushing_2pt_conversions",
            "receiving_2pt_conversions",
        )
    )
    fumbles_lost = sum(
        int(stats.get(field, 0) or 0)
        for field in ("sack_fumbles_lost", "rushing_fumbles_lost", "receiving_fumbles_lost")
    )
    return [
        int(stats.get("passing_yards", 0) or 0),
        int(stats.get("rushing_yards", 0) or 0),
        int(stats.get("receiving_yards", 0) or 0),
        touchdowns,
        two_points,
        int(stats.get("passing_interceptions", 0) or 0),
        int(turnover_tds.get("pick_sixes", 0) or 0),
        fumbles_lost,
        int(turnover_tds.get("fumble_sixes", 0) or 0),
        0,
    ]


def made_field_goal_distances(data, stats: dict[str, Any]) -> list[int]:
    player_id = stats.get("player_id")
    distances: list[int] = []
    if player_id:
        made = data.pbp.filter(
            (pl.col("kicker_player_id") == player_id)
            & (pl.col("field_goal_attempt") == 1)
            & (pl.col("field_goal_result") == "made")
        )
        if made.height:
            distances = [int(value) for value in made["kick_distance"].drop_nulls().to_list()]

    bucket_distances = (
        [20] * int((stats.get("fg_made_0_19", 0) or 0) + (stats.get("fg_made_20_29", 0) or 0))
        + [35] * int(stats.get("fg_made_30_39", 0) or 0)
        + [45] * int(stats.get("fg_made_40_49", 0) or 0)
        + [55] * int(stats.get("fg_made_50_59", 0) or 0)
        + [60] * int(stats.get("fg_made_60_", 0) or 0)
    )
    if len(distances) != len(bucket_distances):
        distances = bucket_distances
    if len(distances) > 7:
        raise ValueError(f"More than seven made field goals for {stats.get('player_display_name')}")
    return distances


def kicker_raw(data, stats: dict[str, Any] | None) -> list[int]:
    if stats is None:
        return [0] * 10
    distances = made_field_goal_distances(data, stats)
    fg_missed = int(stats.get("fg_missed", 0) or 0) + int(stats.get("fg_blocked", 0) or 0)
    pat_missed = int(stats.get("pat_missed", 0) or 0) + int(stats.get("pat_blocked", 0) or 0)
    return (
        distances
        + [0] * (7 - len(distances))
        + [int(stats.get("pat_made", 0) or 0), fg_missed, pat_missed]
    )


def defense_raw(data, team: str) -> list[int]:
    game_info = data.get_game_info(team)
    team_stats = data.get_team_stats(team) or {}
    opponent_stats = data.get_opponent_stats(team) or {}
    sack_info = data.get_defensive_sacks(team)
    blocked_punts = data.get_blocked_punts(team)
    blocked_kick_tds = data.get_blocked_kick_tds(team)
    opponent_stats["_blocked_punts"] = blocked_punts
    opponent_stats["_blocked_kick_tds"] = blocked_kick_tds

    our_recoveries = int(team_stats.get("fumble_recovery_opp", 0) or 0)
    opponent_fumbles_lost = sum(
        int(opponent_stats.get(field, 0) or 0)
        for field in ("sack_fumbles_lost", "rushing_fumbles_lost", "receiving_fumbles_lost")
    )
    fumble_recoveries = max(our_recoveries, opponent_fumbles_lost)
    blocked_fg = int(opponent_stats.get("fg_blocked", 0) or 0)
    total_def_tds = sum(
        int(value or 0)
        for value in (
            team_stats.get("def_tds", 0),
            team_stats.get("fumble_recovery_tds", 0),
            blocked_kick_tds,
        )
    )
    return [
        int((game_info or {}).get("points_allowed", 0) or 0),
        int(sack_info["value"]),
        int(team_stats.get("def_interceptions", 0) or 0),
        fumble_recoveries,
        total_def_tds,
        int(team_stats.get("def_safeties", 0) or 0),
        blocked_fg + blocked_punts,
        int(opponent_stats.get("pat_blocked", 0) or 0),
        0,
        0,
    ]


def score_from_raw(position: str, raw: list[float | int | None]) -> float:
    values = [float(value or 0) for value in raw]
    if position in {"QB", "RB", "WR", "TE"}:
        passing_yards, rushing_yards, receiving_yards, touchdowns, two_points = values[:5]
        interceptions, pick_sixes, fumbles, fumble_sixes = values[5:9]
        passing_points = 0 if passing_yards < 200 else 2 + int((passing_yards - 200) // 50)
        if position == "TE":
            rush_threshold, receive_threshold, combined_threshold = 50, 50, 75
        else:
            rush_threshold, receive_threshold, combined_threshold = 75, 75, 100
        rushing_points = (
            0 if rushing_yards < rush_threshold else 2 + int((rushing_yards - rush_threshold) // 25)
        )
        receiving_points = (
            0
            if receiving_yards < receive_threshold
            else 2 + int((receiving_yards - receive_threshold) // 25)
        )
        combined_yards = rushing_yards + receiving_yards
        combined_points = (
            0
            if combined_yards < combined_threshold
            else 2 + int((combined_yards - combined_threshold) // 25)
        )
        yardage_points = max(rushing_points + receiving_points, combined_points)
        return max(
            0,
            passing_points
            + yardage_points
            + 6 * touchdowns
            + 2 * two_points
            - interceptions
            - 2 * pick_sixes
            - fumbles
            - 2 * fumble_sixes,
        )
    if position == "K":
        field_goals = values[:7]
        pat_made, fg_missed, pat_missed = values[7:10]
        made_points = sum(
            0 if distance < 1 else 1 if distance < 30 else 2 if distance < 40 else 3 if distance < 50 else 4
            for distance in field_goals
        )
        return max(0, made_points + pat_made - 2 * fg_missed - pat_missed)
    if position == "DF":
        points_allowed, sacks, interceptions, fumbles, touchdowns, safeties, blocks, blocked_pats = values[:8]
        if points_allowed == 0:
            allowed_points = 8
        elif points_allowed <= 9:
            allowed_points = 6
        elif points_allowed <= 13:
            allowed_points = 4
        elif points_allowed <= 17:
            allowed_points = 2
        elif points_allowed <= 27:
            allowed_points = 0
        elif points_allowed <= 31:
            allowed_points = -2
        elif points_allowed <= 35:
            allowed_points = -4
        else:
            allowed_points = -6
        return max(
            0,
            allowed_points
            + sacks
            + 2 * (interceptions + fumbles + safeties + blocks)
            + 4 * touchdowns
            + blocked_pats,
        )
    if position == "HC":
        return values[0]
    raise ValueError(position)


def score_starters(teams: list[FantasyTeam], scorer: OPFLScorer) -> None:
    data = scorer.data
    for team in teams:
        for starter in team.starters:
            if not game_complete(data, starter.nfl_team):
                starter.pending = True
                starter.score = None
                starter.raw_inputs = [None] * 10
                continue

            result = scorer.score_player(starter.name, starter.nfl_team, starter.position)
            starter.score = float(result.total_points)
            starter.found = result.found_in_stats
            starter.matched_name = result.matched_name

            if starter.position in {"QB", "RB", "WR", "TE"}:
                stats = player_stats(data, starter)
                starter.raw_inputs = offensive_raw(data, starter, stats)
            elif starter.position == "K":
                stats = player_stats(data, starter)
                starter.raw_inputs = kicker_raw(data, stats)
            elif starter.position == "DF":
                starter.raw_inputs = defense_raw(data, starter.nfl_team)
            else:
                starter.raw_inputs = [excel_number(starter.score)] + [None] * 9

            raw_score = score_from_raw(starter.position, starter.raw_inputs)
            if abs(raw_score - starter.score) > 1e-9:
                raise ValueError(
                    f"Raw inputs do not reconcile for {starter.name}: {raw_score} vs {starter.score}"
                )


def populate_scoring(workbook, teams: list[FantasyTeam]) -> None:
    sheet = workbook["Scoring"]
    for team in teams:
        block_start = 2 + 12 * team.index
        if len(team.starters) != len(SCORING_OFFSETS):
            raise ValueError(f"Unexpected starter count for {team.display_name}")
        for starter, offset in zip(team.starters, SCORING_OFFSETS):
            row = block_start + offset
            for column, value in enumerate(starter.raw_inputs, start=3):
                sheet.cell(row, column).value = value


def make_week_sheet(workbook, teams: list[FantasyTeam]) -> None:
    if "W2" in workbook.sheetnames:
        del workbook["W2"]
    week_sheet = workbook.copy_worksheet(workbook["W1"])
    week_sheet.title = "W2"
    roster = workbook["Rosters"]
    matchups = workbook["Matchups"]

    for row in range(1, roster.max_row + 1):
        for column in range(1, 20):
            source = roster.cell(row, column)
            value = None if isinstance(source.value, str) and source.value.startswith("=") else source.value
            copy_cell(source, week_sheet.cell(row, column), value)

    for column_letter, dimension in roster.column_dimensions.items():
        if openpyxl.utils.column_index_from_string(column_letter) <= 19:
            week_sheet.column_dimensions[column_letter].width = dimension.width
            week_sheet.column_dimensions[column_letter].hidden = dimension.hidden
    for row, dimension in roster.row_dimensions.items():
        week_sheet.row_dimensions[row].height = dimension.height
        week_sheet.row_dimensions[row].hidden = dimension.hidden

    for team in teams:
        points_col = team.player_col - 2
        for starter in team.starters:
            week_sheet.cell(starter.roster_row, points_col).value = excel_number(starter.score)
        total_row = 29 if team.header_row == 1 else 67
        week_sheet.cell(total_row, team.player_col).value = excel_number(team.total)

    for row in range(1, matchups.max_row + 1):
        for source_column in range(1, 6):
            source = matchups.cell(row, source_column)
            target_column = 21 + source_column
            value = None if isinstance(source.value, str) and source.value.startswith("=") else source.value
            copy_cell(source, week_sheet.cell(row, target_column), value)

    code_to_team = {
        1: teams[7],
        2: teams[2],
        3: teams[8],
        4: teams[1],
        5: teams[5],
        6: teams[10],
        7: teams[3],
        8: teams[9],
        9: teams[0],
        10: teams[6],
        11: teams[11],
        12: teams[4],
    }
    for matchup_index, matchup_row in enumerate(range(2, 8)):
        left = code_to_team[int(matchups.cell(matchup_row, 8).value)]
        right = code_to_team[int(matchups.cell(matchup_row, 9).value)]
        start_row = 1 + 12 * matchup_index
        week_sheet.cell(start_row, 22).value = left.display_name
        week_sheet.cell(start_row, 25).value = right.display_name
        for lineup_index, (left_starter, right_starter) in enumerate(
            zip(left.starters, right.starters), start=1
        ):
            week_sheet.cell(start_row + lineup_index, 22).value = left_starter.roster_value
            week_sheet.cell(start_row + lineup_index, 23).value = excel_number(left_starter.score)
            week_sheet.cell(start_row + lineup_index, 24).value = None
            week_sheet.cell(start_row + lineup_index, 25).value = right_starter.roster_value
            week_sheet.cell(start_row + lineup_index, 26).value = excel_number(right_starter.score)
        week_sheet.cell(start_row + 10, 23).value = excel_number(left.total)
        week_sheet.cell(start_row + 10, 26).value = excel_number(right.total)

    workbook._sheets.remove(week_sheet)
    w1_index = workbook._sheets.index(workbook["W1"])
    workbook._sheets.insert(w1_index + 1, week_sheet)
    active_index = workbook._sheets.index(week_sheet)
    for sheet in workbook.worksheets:
        sheet.sheet_view.tabSelected = False
    week_sheet.sheet_view.tabSelected = True
    workbook._active_sheet_index = active_index
    if workbook.views:
        workbook.views[0].activeTab = active_index
        workbook.views[0].firstSheet = active_index


def main() -> None:
    workbook = openpyxl.load_workbook(SOURCE)
    teams = read_teams(workbook)
    scorer = OPFLScorer(SEASON, WEEK)
    score_starters(teams, scorer)
    populate_scoring(workbook, teams)
    make_week_sheet(workbook, teams)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    workbook.close()

    summary = {
        "output": str(OUTPUT.resolve()),
        "teams": [
            {
                "team": team.display_name,
                "total": excel_number(team.total),
                "pending_starters": team.pending_count,
                "starters": [
                    {
                        "position": starter.position,
                        "player": starter.roster_value,
                        "score": excel_number(starter.score),
                        "pending": starter.pending,
                        "found": starter.found,
                        "matched_name": starter.matched_name,
                    }
                    for starter in team.starters
                ],
            }
            for team in teams
        ],
    }
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
