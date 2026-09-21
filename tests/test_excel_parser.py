"""Workbook parsing: player cells, filler rejection, position blocks, team names.

These cover the parsing quirks of the real OPFL workbook — misspelled team
codes, padding zeros, phone numbers sitting in roster columns, and the dozen
different ways the same franchise is spelled across tabs.
"""

import openpyxl
import pytest

from opfl.constants import ALL_TEAM_CODES, resolve_team_code
from opfl.excel_parser import (
    find_position_rows,
    is_valid_player_name,
    parse_player_name,
)


class TestParsePlayerName:
    def test_splits_name_and_team(self):
        assert parse_player_name('Patrick Mahomes (KC)') == ('Patrick Mahomes', 'KC')

    def test_uppercases_mixed_case_team_codes(self):
        """The workbook writes 'Chi', not 'CHI'."""
        assert parse_player_name('Caleb Williams (Chi)') == ('Caleb Williams', 'CHI')

    def test_normalizes_opfl_specific_team_codes(self):
        """OPFL uses LAR/ARZ/JAC where nflverse uses LA/ARI/JAX."""
        assert parse_player_name('Puka Nacua (LAR)') == ('Puka Nacua', 'LA')
        assert parse_player_name('Jacoby Brissett (Arz)') == ('Jacoby Brissett', 'ARI')
        assert parse_player_name('Trevor Lawrence (JAC)') == ('Trevor Lawrence', 'JAX')

    def test_defense_is_a_bare_team_name(self):
        assert parse_player_name('Baltimore') == ('Baltimore', 'BAL')
        assert parse_player_name('LA Rams') == ('LA Rams', 'LA')
        assert parse_player_name('NY Giants') == ('NY Giants', 'NYG')

    def test_trailing_whitespace_is_stripped(self):
        """The workbook has 'Philadelphia ' with a trailing space."""
        assert parse_player_name('Philadelphia ') == ('Philadelphia', 'PHI')

    def test_apostrophes_and_punctuation_survive(self):
        assert parse_player_name("Ja'Marr Chase (Cin)") == ("Ja'Marr Chase", 'CIN')
        assert parse_player_name('Amon-Ra St. Brown (Det)') == ('Amon-Ra St. Brown', 'DET')

    def test_name_without_a_team_is_returned_as_is(self):
        assert parse_player_name('Zach Ertz ()') == ('Zach Ertz ()', '')

    def test_empty_input(self):
        assert parse_player_name('') == ('', '')
        assert parse_player_name(None) == ('', '')


class TestIsValidPlayerName:
    def test_accepts_real_names(self):
        assert is_valid_player_name('Josh Allen')
        assert is_valid_player_name('Denver')
        assert is_valid_player_name('Ben Johnson', 'HC')

    @pytest.mark.parametrize('filler', ['0', '0.0', '', '   ', None])
    def test_rejects_padding_zeros_and_blanks(self, filler):
        """Roster blocks are padded with zeros that must not become players."""
        assert not is_valid_player_name(filler)

    @pytest.mark.parametrize(
        'phone',
        ['407-5858 cell', '325-1289 (J)', '(925) 518-5773 cell', '913-3136 cell'],
    )
    def test_rejects_phone_numbers(self, phone):
        """The PH block sits in the same columns as the roster."""
        assert not is_valid_player_name(phone)

    def test_rejects_names_with_no_letters(self):
        assert not is_valid_player_name('---')

    def test_head_coach_names_must_start_with_a_letter(self):
        assert not is_valid_player_name('1st Choice', 'HC')
        assert is_valid_player_name('Dan Campbell', 'HC')


class TestFindPositionRows:
    def _sheet(self, column_a):
        """Build a worksheet with the given column-A labels and filler in column D."""
        wb = openpyxl.Workbook()
        ws = wb.active
        for row, label in enumerate(column_a, start=1):
            if label is not None:
                ws.cell(row=row, column=1).value = label
            ws.cell(row=row, column=4).value = 'filler'
        return ws

    def test_groups_rows_under_their_position_label(self):
        ws = self._sheet(['QB', None, None, 'RB', None])
        rows = find_position_rows(ws, 1, 5)
        assert rows['QB'] == [1, 2, 3]
        assert rows['RB'] == [4, 5]

    def test_non_position_label_ends_the_roster_section(self):
        """PH (phone numbers) and TS (taxi squad) follow each roster and are not lineup rows."""
        ws = self._sheet(['HC', None, 'PH', None, 'TS', None])
        rows = find_position_rows(ws, 1, 6)
        assert rows['HC'] == [1, 2]
        assert 'PH' not in rows
        assert 'TS' not in rows

    def test_whitespace_only_label_does_not_end_the_section(self):
        """Column A of the second roster block holds '   ', not an empty cell."""
        ws = self._sheet(['QB', '   ', None])
        rows = find_position_rows(ws, 1, 3)
        assert rows['QB'] == [1, 2, 3]


class TestResolveTeamCode:
    def test_every_canonical_owner_resolves(self):
        for code in ALL_TEAM_CODES:
            assert resolve_team_code(code) == code

    @pytest.mark.parametrize(
        ('spelling', 'code'),
        [
            # Rosters tab headers
            ('KIRK/DAVID (1)', 'K/D'),
            ('LAM (7)', 'STL'),
            ('KEMP/A/M (6)', 'KAM'),
            ('GREG/GRIFFIN (3)', 'G/G'),
            # Matchups tab block headers
            ('Kirk/D', 'K/D'),
            ('Steve L', 'STL'),
            ('K/A/M', 'KAM'),
            ('Greg/G', 'G/G'),
            ('Wes/B', 'W/B'),
            # Standings tab
            ('Bill/Wes', 'W/B'),
            ('Greg/Griff', 'G/G'),
            ('Jarrett/M', 'J/M'),
            # Schedule key column
            ('SL', 'STL'),
            ('KL', 'KEV'),
            ('w/b', 'W/B'),
            ('andrew', 'AND'),
        ],
    )
    def test_workbook_spellings_all_resolve(self, spelling, code):
        assert resolve_team_code(spelling) == code

    def test_unknown_name_returns_empty_rather_than_guessing(self):
        """Callers flag an unresolved name instead of silently inventing a team."""
        assert resolve_team_code('Somebody Else Entirely') == ''
        assert resolve_team_code('') == ''
        assert resolve_team_code(None) == ''
