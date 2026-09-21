"""Excel-to-JSON ingestion helpers for OPFL season workbooks."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .excel_parser import parse_player_name
from .league import POSITIONS, season_config, starter_count

TEAM_COLUMNS = (4, 7, 10, 13, 16, 19)
BLOCKS = ((1, 1, 38), (39, 39, 80))
HEADER_PATTERN = re.compile(r"^(.+?)\s*\(\d+\)$")
WEEKLY_WORKBOOK_DIRECTORY = "workbooks"


class ImportValidationError(ValueError):
    """Raised when an Excel import cannot meet the JSON data contract."""


def load_teams(path: str | Path) -> list[dict[str, str]]:
    with Path(path).open(encoding="utf-8") as handle:
        teams = json.load(handle)["teams"]
    if len(teams) != 12:
        raise ImportValidationError(f"Expected 12 teams, found {len(teams)}")
    return teams


def team_aliases(teams: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    aliases: dict[str, dict[str, str]] = {}
    for team in teams:
        names = {team["name"], team["owner"], team["abbrev"]}
        for name in names:
            aliases[_normalize_owner(name)] = team
    aliases.update(
        {
            "GREG/GRIFF": aliases["GREG/GRIFFIN"],
            "GREG/G": aliases["GREG/GRIFFIN"],
            "JARRETT/M": aliases["JARRETT/MATT"],
            "KIRK/D": aliases["KIRK/DAVID"],
            "K/A/M": aliases["KEMP/A/M"],
            "ERIC/J": aliases["ERIC/JEFF"],
            "WES/B": aliases["WES/BILL"],
            "STEVE L": aliases["STEVE L."],
            "LAM": aliases["STEVE L."],
        }
    )
    return aliases


def make_player_id(position: str, name: str, nfl_team: str) -> str:
    raw = f"{position}|{name.casefold()}|{nfl_team.upper()}"
    digest = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:10]
    slug = re.sub(r"[^a-z0-9]+", "-", name.casefold()).strip("-")[:36]
    return f"{position.lower()}-{slug}-{digest}"


def weekly_workbook_path(data_root: str | Path, season: int, week: int) -> Path:
    """Return the required source workbook path for one season/week."""
    if week < 1 or week > 18:
        raise ImportValidationError(f"Week must be between 1 and 18, found {week}")
    return (
        Path(data_root) / "seasons" / str(season) / WEEKLY_WORKBOOK_DIRECTORY / f"week_{week}.xlsx"
    )


def extract_sheet(
    worksheet: Worksheet,
    teams: list[dict[str, str]],
    season: int,
    week: int | None = None,
) -> list[dict[str, Any]]:
    aliases = team_aliases(teams)
    extracted: list[dict[str, Any]] = []
    seen: set[str] = set()
    for header_row, start_row, end_row in BLOCKS:
        position_rows = _position_rows(worksheet, start_row, end_row)
        for player_column in TEAM_COLUMNS:
            raw_header = worksheet.cell(header_row, player_column).value
            if not raw_header:
                continue
            match = HEADER_PATTERN.match(str(raw_header).strip())
            owner = match.group(1).strip() if match else str(raw_header).strip()
            team = aliases.get(_normalize_owner(owner))
            if not team:
                raise ImportValidationError(
                    f"Unknown owner {owner!r} in {worksheet.title}!{worksheet.cell(header_row, player_column).coordinate}"
                )
            abbreviation = team["abbrev"]
            if abbreviation in seen:
                raise ImportValidationError(f"Duplicate team {abbreviation} in {worksheet.title}")
            seen.add(abbreviation)
            roster: list[dict[str, Any]] = []
            for position, rows in position_rows.items():
                for row in rows:
                    value = worksheet.cell(row, player_column).value
                    if not _valid_player_value(value):
                        continue
                    name, nfl_team = parse_player_name(str(value).strip())
                    score_value = worksheet.cell(row, player_column - 2).value
                    score = _number(score_value) if week is not None else None
                    starter = worksheet.cell(row, player_column - 1).value == "*"
                    roster.append(
                        {
                            "player_id": make_player_id(position, name, nfl_team),
                            "name": name,
                            "nfl_team": nfl_team,
                            "position": position,
                            "starter": starter,
                            **({"score": score, "breakdown": {}} if week is not None else {}),
                        }
                    )
            if not roster:
                raise ImportValidationError(
                    f"No players found for {abbreviation} in {worksheet.title}"
                )
            if week is not None:
                _fill_vacant_starters(roster, abbreviation, season, week)
            item: dict[str, Any] = {
                "abbrev": abbreviation,
                "name": team["name"],
                "owner": team["owner"],
                "roster": roster,
            }
            if week is not None:
                starter_total = sum(player["score"] for player in roster if player["starter"])
                total_row = 29 if header_row == 1 else 67
                recorded_total = worksheet.cell(total_row, player_column).value
                team_total = (
                    _number(recorded_total) if _is_number(recorded_total) else starter_total
                )
                correction = round(team_total - starter_total, 2)
                item.update(
                    {
                        "starter_count": sum(1 for player in roster if player["starter"]),
                        "starter_total": round(starter_total, 2),
                        "manual_correction": correction,
                        "total_score": round(team_total, 2),
                    }
                )
            extracted.append(item)
    if len(extracted) != 12 or seen != {team["abbrev"] for team in teams}:
        raise ImportValidationError(
            f"{worksheet.title} must contain all 12 teams; found {len(extracted)}"
        )
    if week is not None:
        _validate_week_lineups(extracted, season)
    return extracted


def import_workbook(
    excel_path: str | Path,
    teams: list[dict[str, str]],
    season: int,
) -> tuple[dict[str, Any], dict[int, dict[str, Any]], dict[int, dict[str, Any]]]:
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    try:
        if "Rosters" not in workbook.sheetnames:
            raise ImportValidationError("Workbook must contain a Rosters sheet")
        current = extract_sheet(workbook["Rosters"], teams, season)
        rosters = {
            "season": season,
            "teams": {
                item["abbrev"]: {
                    "name": item["name"],
                    "owner": item["owner"],
                    "players": [
                        {key: value for key, value in player.items() if key != "starter"}
                        for player in item["roster"]
                    ],
                    "taxi": [],
                }
                for item in current
            },
        }
        weeks: dict[int, dict[str, Any]] = {}
        lineups: dict[int, dict[str, Any]] = {}
        week_sheets = sorted(
            (
                (int(match.group(1)), sheet)
                for sheet in workbook.sheetnames
                if (match := re.fullmatch(r"W(\d+)", sheet))
            ),
            key=lambda item: item[0],
        )
        for week, sheet in week_sheets:
            week_teams = extract_sheet(workbook[sheet], teams, season, week)
            ranked = sorted(week_teams, key=lambda item: (-item["total_score"], item["abbrev"]))
            ranks = {item["abbrev"]: rank for rank, item in enumerate(ranked, 1)}
            for item in week_teams:
                item["score_rank"] = ranks[item["abbrev"]]
            weeks[week] = {
                "season": season,
                "week": week,
                "phase": "regular" if week <= 15 else "playoffs",
                "status": "final",
                "score_source": "excel_archive" if season == 2025 else "json_scorer",
                "teams": week_teams,
                "matchups": [],
            }
            lineups[week] = {
                "season": season,
                "week": week,
                "lineups": {
                    item["abbrev"]: [
                        player["player_id"] for player in item["roster"] if player["starter"]
                    ]
                    for item in week_teams
                },
                "manual_corrections": {
                    item["abbrev"]: item["manual_correction"]
                    for item in week_teams
                    if item["manual_correction"]
                },
            }
        return rosters, weeks, lineups
    finally:
        workbook.close()


def import_weekly_workbook(
    excel_path: str | Path,
    teams: list[dict[str, str]],
    season: int,
    week: int,
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Import one 2026-style workbook containing Rosters and Matchups tabs."""
    path = Path(excel_path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        required_sheets = {"Rosters", "Matchups"}
        missing_sheets = required_sheets - set(workbook.sheetnames)
        if missing_sheets:
            raise ImportValidationError(
                f"{path.name} is missing sheet(s): {', '.join(sorted(missing_sheets))}"
            )
        current = extract_sheet(workbook["Rosters"], teams, season)
        _validate_week_lineups(current, season)
        pairings = _extract_matchup_pairings(workbook["Matchups"], teams)
    finally:
        workbook.close()

    roster_teams = {
        item["abbrev"]: {
            "name": item["name"],
            "owner": item["owner"],
            "players": [
                {key: value for key, value in player.items() if key != "starter"}
                for player in item["roster"]
            ],
            "taxi": [],
        }
        for item in current
    }
    rosters = {"season": season, "teams": roster_teams}
    lineups = {
        "season": season,
        "week": week,
        "source_workbook": f"{WEEKLY_WORKBOOK_DIRECTORY}/week_{week}.xlsx",
        "pairings": pairings,
        "lineups": {
            item["abbrev"]: [player["player_id"] for player in item["roster"] if player["starter"]]
            for item in current
        },
        "roster_snapshot": rosters,
        "manual_corrections": {},
    }
    return rosters, lineups


def attach_regular_matchups(
    weeks: dict[int, dict[str, Any]], schedule: list[dict[str, Any]]
) -> None:
    schedule_by_week = {int(item["week"]): item["matchups"] for item in schedule}
    for week, week_data in weeks.items():
        if week > 15:
            continue
        scores = {team["abbrev"]: team["total_score"] for team in week_data["teams"]}
        week_data["matchups"] = [
            {
                **matchup,
                "away_score": scores[matchup["away"]],
                "home_score": scores[matchup["home"]],
            }
            for matchup in schedule_by_week[week]
        ]


def _extract_matchup_pairings(
    worksheet: Worksheet, teams: list[dict[str, str]]
) -> list[dict[str, str]]:
    aliases = team_aliases(teams)
    pairings: list[dict[str, str]] = []
    seen: set[str] = set()
    row = 1
    while row <= worksheet.max_row:
        left_value = worksheet.cell(row, 1).value
        right_value = worksheet.cell(row, 4).value
        left = aliases.get(_normalize_owner(str(left_value))) if left_value else None
        right = aliases.get(_normalize_owner(str(right_value))) if right_value else None
        if not left or not right:
            row += 1
            continue
        away, home = left["abbrev"], right["abbrev"]
        if away in seen or home in seen:
            raise ImportValidationError(f"Duplicate matchup team in {worksheet.title} row {row}")
        pairings.append({"away": away, "home": home})
        seen.update((away, home))
        row += 12
    expected = {team["abbrev"] for team in teams}
    if len(pairings) != 6 or seen != expected:
        raise ImportValidationError(
            f"{worksheet.title} must contain six matchups covering all 12 teams"
        )
    return pairings


def _position_rows(worksheet: Worksheet, start_row: int, end_row: int) -> dict[str, list[int]]:
    positions: dict[str, list[int]] = {}
    active: str | None = None
    for row in range(start_row, min(end_row, worksheet.max_row) + 1):
        label_value = worksheet.cell(row, 1).value
        label = str(label_value).strip().upper() if label_value is not None else ""
        if label in POSITIONS:
            active = label
            positions.setdefault(active, []).append(row)
            continue
        if label:
            active = None
            continue
        if active and any(worksheet.cell(row, column).value for column in TEAM_COLUMNS):
            positions[active].append(row)
    return positions


def _valid_player_value(value: object) -> bool:
    return value is not None and bool(str(value).strip()) and not _is_number(value)


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _number(value: object) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError) as exc:
        raise ImportValidationError(f"Expected a numeric score, found {value!r}") from exc


def _normalize_owner(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().upper())


def _validate_week_lineups(teams: list[dict[str, Any]], season: int) -> None:
    config = season_config(season)
    expected = starter_count(config)
    base_counts = {
        slot["name"]: int(slot["count"])
        for slot in config["lineup_slots"]
        if slot["name"] != "FLEX" and int(slot["count"]) > 0
    }
    flex_count = next(
        (int(slot["count"]) for slot in config["lineup_slots"] if slot["name"] == "FLEX"),
        0,
    )
    flex_eligible = next(
        (
            set(slot["eligible_positions"])
            for slot in config["lineup_slots"]
            if slot["name"] == "FLEX"
        ),
        set(),
    )
    for team in teams:
        starters = [player for player in team["roster"] if player["starter"]]
        if len(starters) != expected:
            raise ImportValidationError(
                f"{team['abbrev']} has {len(starters)} starters; expected {expected}"
            )
        counts: dict[str, int] = {position: 0 for position in POSITIONS}
        for player in starters:
            counts[player["position"]] += 1
        flex_used = 0
        for position, count in counts.items():
            overflow = count - base_counts.get(position, 0)
            if overflow > 0:
                if position not in flex_eligible:
                    raise ImportValidationError(
                        f"{team['abbrev']} has too many {position} starters"
                    )
                flex_used += overflow
        if flex_used != flex_count:
            raise ImportValidationError(
                f"{team['abbrev']} uses {flex_used} FLEX starters; expected {flex_count}"
            )


def _fill_vacant_starters(
    roster: list[dict[str, Any]], abbreviation: str, season: int, week: int
) -> None:
    """Represent an officially unfilled lineup slot as an explicit zero-point starter."""
    config = season_config(season)
    required = {
        slot["name"]: int(slot["count"])
        for slot in config["lineup_slots"]
        if slot["name"] != "FLEX" and int(slot["count"]) > 0
    }
    counts = {position: 0 for position in POSITIONS}
    for player in roster:
        if player["starter"]:
            counts[player["position"]] += 1
    for position, count in required.items():
        missing = count - counts[position]
        for index in range(max(0, missing)):
            roster.append(
                {
                    "player_id": (
                        f"vacant-{season}-{week}-{re.sub(r'[^a-z0-9]+', '-', abbreviation.casefold()).strip('-')}"
                        f"-{position.casefold()}-{index + 1}"
                    ),
                    "name": "No starter submitted",
                    "nfl_team": "",
                    "position": position,
                    "starter": True,
                    "score": 0.0,
                    "breakdown": {},
                    "placeholder": True,
                }
            )
